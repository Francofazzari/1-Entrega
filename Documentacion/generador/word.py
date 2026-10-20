# -*- coding: utf-8 -*-
"""Piezas para armar el Word con python-docx: estilos, campos (indice, epigrafes, referencias), tablas y figuras.

La base es la entrega anterior del alumno (copia de solo lectura): de ella se toman el encabezado y el pie
oficiales de la UAI y los estilos; el cuerpo se arma de cero.
"""
import re
from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor

CENTRO = WD_ALIGN_PARAGRAPH.CENTER
JUSTIFICADO = WD_ALIGN_PARAGRAPH.JUSTIFY

AZUL = '1F4E79'       # mismos colores que los diagramas
AZUL_CLARO = 'EAF1F8'
GRIS = '8A99AD'
TEXTO = '1F2937'
OSCURO = '2D3E50'

ANCHO = 15.5          # ancho de texto (cm): A4 con margenes de 2,75 cm
ALTO_FIG = 19.0       # alto maximo de una figura sola en su pagina (cm), con su epigrafe
ALTO_TRAS_TITULO = 16.5   # alto maximo de la primera figura de una seccion (debajo del titulo y su introduccion)
ESCALA_MAX = 0.030    # cm por px de diagrama (0,30 mm/px): los diagramas chicos no se agrandan de mas
ROTAR_SI = 1.15       # se gira 90 grados si asi la letra queda al menos 15 % mas grande

# orden de los hijos de tcPr y tblPr segun el esquema de WordprocessingML
_ORDEN_TCPR = ['cnfStyle', 'tcW', 'gridSpan', 'hMerge', 'vMerge', 'tcBorders', 'shd', 'noWrap', 'tcMar',
               'textDirection', 'tcFitText', 'vAlign', 'hideMark']
_ORDEN_TBLPR = ['tblStyle', 'tblpPr', 'tblOverlap', 'bidiVisual', 'tblStyleRowBandSize', 'tblStyleColBandSize',
                'tblW', 'jc', 'tblCellSpacing', 'tblInd', 'tblBorders', 'shd', 'tblLayout', 'tblCellMar', 'tblLook']


def _insertar(padre, nuevo, orden):
    """Inserta 'nuevo' en 'padre' respetando el orden del esquema (reemplaza si ya habia uno igual)."""
    nombre = nuevo.tag.split('}')[1]
    viejo = padre.find(qn('w:' + nombre))
    if viejo is not None:
        padre.replace(viejo, nuevo)
        return nuevo
    despues = orden[orden.index(nombre) + 1:]
    for hijo in padre:
        if hijo.tag.split('}')[1] in despues:
            hijo.addprevious(nuevo)
            return nuevo
    padre.append(nuevo)
    return nuevo


def _el(nombre, **attrs):
    e = OxmlElement('w:' + nombre)
    for k, v in attrs.items():
        e.set(qn('w:' + k), str(v))
    return e


class Word:
    def __init__(self, ref, logo, carpeta_rotadas):
        self.doc = Document(str(ref))
        self.logo = Path(logo)
        self.rotadas = Path(carpeta_rotadas)
        self.n_fig = 0
        self.n_marcador = 1000
        self.figuras = []          # (numero, figura, puntos, girada)
        body = self.doc.element.body
        for el in list(body):
            body.remove(el)
        self._estilos()

    # ------------------------------------------------------------------ estilos
    def estilo(self, nombre_interno):
        for s in self.doc.styles:
            if s.element.name_val == nombre_interno:
                return s
        return None

    def _fuente(self, estilo, nombre='Arial', tam=None, negrita=None, color=None, cursiva=None):
        f = estilo.font
        f.name = nombre
        rpr = estilo.element.get_or_add_rPr()
        rf = rpr.get_or_add_rFonts()
        for k in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
            # una fuente del tema (por ejemplo, Cambria en el titulo del indice) le gana a la fuente explicita
            if rf.get(qn('w:' + k)) is not None:
                del rf.attrib[qn('w:' + k)]
        for k in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
            rf.set(qn('w:' + k), nombre)
        if tam:
            f.size = Pt(tam)
        if negrita is not None:
            f.bold = negrita
        if cursiva is not None:
            f.italic = cursiva
        if color:
            f.color.rgb = RGBColor.from_string(color)

    def _estilos(self):
        st = self.doc.styles
        normal = st['Normal']
        # Normal no se toca: lo usan los cuadros de texto del encabezado y del pie oficiales, que tienen alto fijo.
        # El cuerpo usa su propio estilo.
        texto = st.add_style('Texto', WD_STYLE_TYPE.PARAGRAPH)
        texto.base_style = normal
        texto.quick_style = True
        self._fuente(texto, tam=11)
        pf = texto.paragraph_format
        pf.space_after = Pt(6)
        pf.line_spacing = 1.15
        # titulos: Arial 14 (titulos), 13 (subtitulos) y 12; el cuerpo en 11 (pauta de la catedra: 12 como maximo)
        for nivel, tam, antes, despues in ((1, 14, 0, 16), (2, 13, 18, 10), (3, 12, 14, 8), (4, 11, 10, 6)):
            s = self.estilo('heading %d' % nivel)
            self._fuente(s, tam=tam, negrita=True, color=AZUL)
            p = s.paragraph_format
            p.space_before = Pt(antes)
            p.space_after = Pt(despues)
            p.keep_with_next = True
            p.left_indent = Cm(0)
            p.line_spacing = 1.0
            if nivel == 1:
                p.page_break_before = True
        cap = self.estilo('caption')
        if cap is None:
            cap = st.add_style('Caption', WD_STYLE_TYPE.PARAGRAPH)
            cap.base_style = normal
        self._fuente(cap, tam=9.5, negrita=False, color=OSCURO, cursiva=False)
        p = cap.paragraph_format
        p.alignment = CENTRO
        p.space_before = Pt(4)
        p.space_after = Pt(18)
        p.keep_with_next = False
        p.line_spacing = 1.0
        for i, sang in ((1, 0), (2, 0.5), (3, 1.0)):
            s = self.estilo('toc %d' % i)
            self._fuente(s, tam=10, negrita=(i == 1))
            s.paragraph_format.left_indent = Cm(sang)
            s.paragraph_format.space_before = Pt(4 if i == 1 else 1)
            s.paragraph_format.space_after = Pt(1)
            s.paragraph_format.line_spacing = 1.0
        tf = self.estilo('table of figures')
        if tf is None:
            tf = st.add_style('table of figures', WD_STYLE_TYPE.PARAGRAPH)
            tf.base_style = normal
        self._fuente(tf, tam=9.5)
        tf.paragraph_format.left_indent = Cm(0)
        tf.paragraph_format.space_after = Pt(2)
        tf.paragraph_format.line_spacing = 1.0
        th = self.estilo('TOC Heading')
        if th is not None:
            self._fuente(th, tam=14, negrita=True, color=AZUL)
            th.paragraph_format.space_after = Pt(14)

    # ------------------------------------------------------------------ encabezado oficial
    def encabezado(self, rid, reemplazos, achicar=()):
        """Completa los textos del encabezado oficial (cuadros de texto de la plantilla de la UAI).

        achicar: textos originales cuyo reemplazo va en 8 pt porque no entra en su cuadro (por ejemplo, la fecha).
        """
        parte = self.doc.part.related_parts[rid]
        hechos = set()
        for p in parte._element.iter(qn('w:p')):
            ts = list(p.iter(qn('w:t')))
            texto = ''.join(t.text or '' for t in ts)
            if texto in reemplazos:
                ts[0].text = reemplazos[texto]
                ts[0].set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
                for t in ts[1:]:
                    t.text = ''
                hechos.add(texto)
                if texto in achicar:
                    for r in p.findall(qn('w:r')):
                        rpr = r.find(qn('w:rPr'))
                        if rpr is None:
                            rpr = OxmlElement('w:rPr')
                            r.insert(0, rpr)
                        for nombre in ('sz', 'szCs'):
                            viejo = rpr.find(qn('w:' + nombre))
                            if viejo is not None:
                                rpr.remove(viejo)
                            rpr.append(_el(nombre, val=16))
        return set(reemplazos) - hechos

    def pie_centrado(self, rid, ancho_pt=120):
        """El numero de pagina del pie oficial esta en un cuadro de texto de 70 pt: con 'Pagina 125 de 264' no
        entra. Se lo ensancha hacia los dos lados y se centra el texto, asi sigue centrado en la hoja."""
        parte = self.doc.part.related_parts[rid]
        raiz = parte._element
        WP = '{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}'
        A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
        V = '{urn:schemas-microsoft-com:vml}'
        nuevo = int(ancho_pt * 12700)
        for anc in raiz.iter(WP + 'anchor'):
            ext = anc.find(WP + 'extent')
            viejo = int(ext.get('cx'))
            ext.set('cx', str(nuevo))
            pos = anc.find(WP + 'positionH').find(WP + 'posOffset')
            pos.text = str(int(pos.text) - (nuevo - viejo) // 2)
            for e in anc.iter(A + 'ext'):
                e.set('cx', str(nuevo))
        for shp in raiz.iter(V + 'shape'):
            st = shp.get('style') or ''
            ml = re.search(r'margin-left:([\d.]+)pt', st)
            an = re.search(r'width:([\d.]+)pt', st)
            if ml and an:
                izq = float(ml.group(1)) - (ancho_pt - float(an.group(1))) / 2
                st = st.replace(ml.group(0), 'margin-left:%.2fpt' % izq).replace(an.group(0), 'width:%.2fpt' % ancho_pt)
                shp.set('style', st)
        orden_ppr = ['pStyle', 'keepNext', 'keepLines', 'pageBreakBefore', 'framePr', 'widowControl', 'numPr',
                     'suppressLineNumbers', 'pBdr', 'shd', 'tabs', 'suppressAutoHyphens', 'kinsoku', 'wordWrap',
                     'overflowPunct', 'topLinePunct', 'autoSpaceDE', 'autoSpaceDN', 'bidi', 'adjustRightInd',
                     'snapToGrid', 'spacing', 'ind', 'contextualSpacing', 'mirrorIndents', 'suppressOverlap', 'jc',
                     'textDirection', 'textAlignment', 'textboxTightWrap', 'outlineLvl', 'divId', 'cnfStyle', 'rPr',
                     'sectPr', 'pPrChange']
        for txbx in raiz.iter(qn('w:txbxContent')):
            for p in txbx.iter(qn('w:p')):
                ppr = p.find(qn('w:pPr'))
                if ppr is None:
                    ppr = OxmlElement('w:pPr')
                    p.insert(0, ppr)
                ind = ppr.find(qn('w:ind'))
                if ind is not None:
                    ppr.remove(ind)
                _insertar(ppr, _el('jc', val='center'), orden_ppr)

    # ------------------------------------------------------------------ secciones
    def seccion(self, con_encabezado, rid_enc=None, rid_pie=None, top=3420, bottom=1526, inicio=None):
        s = OxmlElement('w:sectPr')
        if con_encabezado:
            s.append(OxmlElement('w:headerReference'))
            s[-1].set(qn('w:type'), 'default')
            s[-1].set(qn('r:id'), rid_enc)
            s.append(OxmlElement('w:footerReference'))
            s[-1].set(qn('w:type'), 'default')
            s[-1].set(qn('r:id'), rid_pie)
        s.append(_el('pgSz', w=11910, h=16840))
        s.append(_el('pgMar', top=top, right=1559, bottom=bottom, left=1559, header=710, footer=772, gutter=0))
        if inicio:
            s.append(_el('pgNumType', start=inicio))
        s.append(_el('cols', space=720))
        return s

    def cerrar_seccion(self, sectpr):
        """La seccion actual termina en un parrafo vacio que lleva sus propiedades."""
        par = self.doc.add_paragraph()
        par._p.get_or_add_pPr().append(sectpr)
        par.paragraph_format.space_after = Pt(0)
        return par

    def seccion_final(self, sectpr):
        self.doc.element.body.append(sectpr)

    # ------------------------------------------------------------------ texto
    def _runs(self, par, texto, negrita=False, cursiva=False, tam=None, color=None):
        for parte in re.split(r'(\*\*.+?\*\*)', texto):
            if not parte:
                continue
            b = negrita
            if parte.startswith('**') and parte.endswith('**'):
                parte, b = parte[2:-2], True
            r = par.add_run(parte)
            if b:
                r.bold = True
            if cursiva:
                r.italic = True
            if tam:
                r.font.size = Pt(tam)
            if color:
                r.font.color.rgb = RGBColor.from_string(color)
        return par

    def p(self, texto='', alinear=JUSTIFICADO, negrita=False, cursiva=False, tam=None, color=None, antes=None,
          despues=None, mantener=False, estilo='Texto'):
        par = self.doc.add_paragraph(style=estilo)
        self._runs(par, texto, negrita, cursiva, tam, color)
        if alinear is not None:
            par.alignment = alinear
        pf = par.paragraph_format
        if antes is not None:
            pf.space_before = Pt(antes)
        if despues is not None:
            pf.space_after = Pt(despues)
        if mantener:
            pf.keep_with_next = True
        return par

    def vinetas(self, items, tam=None, sangria=0.9):
        for it in items:
            par = self.doc.add_paragraph(style='Texto')
            pf = par.paragraph_format
            pf.left_indent = Cm(sangria)
            pf.first_line_indent = Cm(-0.45)
            pf.space_after = Pt(3)
            pf.tab_stops.add_tab_stop(Cm(sangria))
            par.alignment = JUSTIFICADO
            self._runs(par, '•\t' + it, tam=tam)

    def titulo(self, nivel, texto, nueva_pagina=False):
        par = self.doc.add_heading(texto, level=nivel)
        if nueva_pagina:
            par.paragraph_format.page_break_before = True
        return par

    def salto_pagina(self):
        par = self.doc.add_paragraph()
        par.add_run().add_break(WD_BREAK.PAGE)
        par.paragraph_format.space_after = Pt(0)
        return par

    # ------------------------------------------------------------------ campos y marcadores
    def campo(self, par, instruccion, resultado='', negrita=False):
        def run(hijo):
            r = OxmlElement('w:r')
            if negrita:
                rpr = OxmlElement('w:rPr')
                rpr.append(OxmlElement('w:b'))
                r.append(rpr)
            r.append(hijo)
            par._p.append(r)
        run(_el('fldChar', fldCharType='begin'))
        it = OxmlElement('w:instrText')
        it.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
        it.text = instruccion
        run(it)
        run(_el('fldChar', fldCharType='separate'))
        t = OxmlElement('w:t')
        t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
        t.text = resultado
        run(t)
        run(_el('fldChar', fldCharType='end'))

    def indice(self, titulo, instruccion, aviso):
        th = self.estilo('TOC Heading')
        par = self.doc.add_paragraph(titulo, style=th) if th is not None else self.p(titulo, negrita=True, tam=14)
        par = self.doc.add_paragraph()
        self.campo(par, instruccion, aviso)
        return par

    # ------------------------------------------------------------------ figuras
    def epigrafe(self, texto):
        self.n_fig += 1
        n = self.n_fig
        par = self.doc.add_paragraph(style=self.estilo('caption'))
        self.n_marcador += 1
        nombre = '_Fig%03d' % n
        par._p.append(_el('bookmarkStart', id=self.n_marcador, name=nombre))
        r = par.add_run('Figura ')
        r.bold = True
        self.campo(par, ' SEQ Figura \\* ARABIC ', str(n), negrita=True)
        par._p.append(_el('bookmarkEnd', id=self.n_marcador))
        r = par.add_run('. ')
        r.bold = True
        par.add_run(texto)
        return n

    def referencia(self, par, n):
        """Campo REF a 'Figura n' (se actualiza solo si la figura cambia de numero)."""
        self.campo(par, ' REF _Fig%03d \\h \\* MERGEFORMAT ' % n, 'Figura %d' % n)

    def figura(self, fig, max_alto=ALTO_FIG, rotar=True, ancho_max=ANCHO, antes=10):
        im = Image.open(fig.ruta)
        wpx, hpx = im.size
        girada = False
        puntos = None
        if fig.tipo == 'proto':
            # capturas de pantalla: como mucho a 96 ppp
            s = min(ancho_max / wpx, 15.0 / hpx, 0.0265)
            ruta, ancho = fig.ruta, wpx * s
        else:
            w, h = wpx / 2.0, hpx / 2.0          # los PNG de los diagramas estan a escala 2
            su = min(ancho_max / w, max_alto / h, ESCALA_MAX)
            sr = min(ancho_max / h, max_alto / w, ESCALA_MAX)
            if rotar and sr > ROTAR_SI * su:
                self.rotadas.mkdir(parents=True, exist_ok=True)
                ruta = self.rotadas / Path(fig.ruta).name
                im.convert('RGB').rotate(90, expand=True).save(ruta)
                s, ancho, girada = sr, h * sr, True
            else:
                ruta, s, ancho = fig.ruta, su, w * su
            if fig.letra:
                puntos = fig.letra * s * 10 / 0.3528
        par = self.doc.add_paragraph()
        par.alignment = CENTRO
        pf = par.paragraph_format
        pf.keep_with_next = True
        pf.space_before = Pt(antes)
        pf.space_after = Pt(2)
        pf.line_spacing = 1.0
        par.add_run().add_picture(str(ruta), width=Cm(ancho))
        n = self.epigrafe(fig.epigrafe)
        fig.numero = n
        self.figuras.append((n, fig, puntos, girada))
        return n

    # ------------------------------------------------------------------ tablas
    def tabla(self, anchos, encabezados=None, filas=(), tam=9.5, centradas=(), repetir=True, alto_min=None):
        t = self.doc.add_table(rows=0, cols=len(anchos))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        tblPr = t._tbl.tblPr
        bordes = OxmlElement('w:tblBorders')
        for lado in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            bordes.append(_el(lado, val='single', sz=4, space=0, color=GRIS))
        _insertar(tblPr, bordes, _ORDEN_TBLPR)
        _insertar(tblPr, _el('tblW', w=int(sum(anchos) * 567), type='dxa'), _ORDEN_TBLPR)
        for gc, w in zip(t._tbl.tblGrid.findall(qn('w:gridCol')), anchos):
            gc.set(qn('w:w'), str(int(w * 567)))
        t._anchos = anchos
        if encabezados:
            fila = self.fila(t, encabezados, tam=tam, negrita=True, color='FFFFFF', fondo=AZUL)
            if repetir:
                fila._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        for f in filas:
            self.fila(t, f, tam=tam, centradas=centradas)
        return t

    def fila(self, t, valores, tam=9.5, negrita=False, color=None, fondo=None, centradas=(), unir=None,
             fondos=None, no_dividir=True, negritas=None, vinetas=None, mantener=False):
        """unir: lista de tramos [(desde, hasta)] de columnas que se combinan; valores va uno por tramo."""
        row = t.add_row()
        if no_dividir:
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        tramos = unir or [(i, i) for i in range(len(t._anchos))]
        for k, (a, b) in enumerate(tramos):
            celda = row.cells[a]
            if b > a:
                celda = celda.merge(row.cells[b])
            ancho = sum(t._anchos[a:b + 1])
            celda.width = Cm(ancho)
            f = (fondos[k] if fondos else None) or fondo
            if f:
                self.fondo(celda, f)
            celda.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            neg = negritas[k] if negritas else negrita
            self.celda(celda, valores[k] if k < len(valores) else '', tam=tam, negrita=neg, color=color,
                       centrada=(k in centradas), vinetas=vinetas)
            if mantener:
                for par in celda.paragraphs:
                    par.paragraph_format.keep_with_next = True
        return row

    def fondo(self, celda, color):
        tcPr = celda._tc.get_or_add_tcPr()
        _insertar(tcPr, _el('shd', val='clear', color='auto', fill=color), _ORDEN_TCPR)

    def celda(self, celda, contenido, tam=9.5, negrita=False, color=None, centrada=False, vinetas=None):
        es_lista = isinstance(contenido, (list, tuple))
        items = contenido if es_lista else [contenido]
        if vinetas is None:
            vinetas = es_lista and len(contenido) > 1
        elif not es_lista:
            vinetas = False         # las viñetas son para las listas; un rótulo o un texto suelto no lleva
        par = celda.paragraphs[0]
        for i, it in enumerate(items):
            if i:
                par = celda.add_paragraph()
            pf = par.paragraph_format
            pf.space_before = Pt(1)
            pf.space_after = Pt(2)
            pf.line_spacing = 1.0
            if centrada:
                par.alignment = CENTRO
            texto = str(it)
            if vinetas:
                pf.left_indent = Cm(0.4)
                pf.first_line_indent = Cm(-0.35)
                pf.tab_stops.add_tab_stop(Cm(0.4))
                texto = '•\t' + texto
            self._runs(par, texto, negrita=negrita, tam=tam, color=color)
        return celda

    # ------------------------------------------------------------------ salida
    def limpiar_relaciones(self):
        """Quita las relaciones del documento base que ya no se usan (encabezados viejos, imagenes)."""
        xml = self.doc.element.xml
        usados = set(re.findall(r'r:(?:id|embed|link)="(rId\d+)"', xml))
        quitadas = []
        for rid, rel in list(self.doc.part.rels.items()):
            tipo = rel.reltype.split('/')[-1]
            if tipo in ('header', 'footer', 'image', 'hyperlink') and rid not in usados:
                self.doc.part.drop_rel(rid)
                quitadas.append(rid)
        return quitadas

    def guardar(self, ruta):
        Path(ruta).parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(str(ruta))
