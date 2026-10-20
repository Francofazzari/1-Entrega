# -*- coding: utf-8 -*-
"""Arma el documento N01-N02 (Word) con todas las figuras del catalogo y las especificaciones de casos de uso.

Uso: python armar_word.py [ruta_salida.docx]

Antes hay que generar los PNG con catalogo.py. El indice y los numeros de las figuras quedan como campos de Word:
actualizar_word.ps1 los actualiza y exporta el PDF.
"""
import sys
from pathlib import Path

from docx.oxml import OxmlElement
from docx.shared import Cm, Pt

import catalogo
import n01
from balanceo import especificaciones, verificar
from comun import (CASOS, NOMBRE_PROCESO, PROCESOS, REQUISITOS, REGLAS, TEXTO_REGLA, SISTEMA_CORTO, AUTOR, FECHA,
                   VERSION, esp)
from n01_diag import ACTIVIDADES
from n02_comun import CORRESPONDENCIA
from word import Word, CENTRO, JUSTIFICADO, AZUL, AZUL_CLARO, ALTO_TRAS_TITULO

SCRATCH = Path(__file__).resolve().parents[2]
REF = SCRATCH / 'ref_entrega.docx'
LOGO = SCRATCH / 'doc' / 'recursos' / 'portada_logo.png'
ROTADAS = SCRATCH / 'doc' / 'png_rot'
SALIDA = Path(r'C:\Users\franco\source\repos\Entrega1software\Documentacion\N01_N02_Agencia_de_Quiniela.docx')

BARRA = 'D6E4F0'          # fondo de las barras de seccion de la plantilla
ANCHOS_PLANTILLA = (1.6, 2.7, 2.0, 3.0, 6.2)


def lista_y(items):
    items = list(items)
    if len(items) <= 1:
        return ''.join(items)
    return ', '.join(items[:-1]) + ' y ' + items[-1]


def cu_nombre(c):
    return '%s %s' % (c, CASOS[c][0])


def requisitos_de(c):
    return [rf for rf, desc, proc, casos in REQUISITOS if c in casos]


# ====================================================================== portada e indices
def portada(w):
    par = w.doc.add_paragraph()
    par.alignment = CENTRO
    par.paragraph_format.space_before = Pt(6)
    par.add_run().add_picture(str(LOGO), width=Cm(12.5))
    w.p('FACULTAD DE TECNOLOGÍA INFORMÁTICA', alinear=CENTRO, negrita=True, tam=12, antes=20, despues=2)
    w.p('TRABAJO DE DIPLOMA', alinear=CENTRO, tam=11, despues=28)
    w.p(SISTEMA_CORTO, alinear=CENTRO, negrita=True, tam=20, color=AZUL, despues=10)
    w.p('N01. Especificación funcional por proceso de negocio', alinear=CENTRO, tam=14, color=AZUL, despues=2)
    w.p('N02. Especificaciones de casos de uso', alinear=CENTRO, tam=14, color=AZUL, despues=10)
    w.p('Gestión Comercial de Quiniela   |   Versión %s   |   %s' % (VERSION, FECHA), alinear=CENTRO, tam=11,
        despues=34)
    for linea in ('Alumno: Franco Tomas Fazzari', 'Localización: Castelar',
                  'Docentes: Julián Rodriguez Escobedo - Horacio Valdez', 'Turno: Noche', 'Año de cursada: 2026',
                  'Comisión: 3A'):
        w.p(linea, alinear=None, tam=12, despues=8)
    w.cerrar_seccion(w.seccion(False, top=1420, bottom=1134))


def indices(w):
    w.indice('Índice', ' TOC \\o "1-3" \\h \\z \\u ',
             'Índice: se completa al actualizar los campos del documento (F9).')
    w.salto_pagina()
    w.indice('Índice de figuras', ' TOC \\h \\z \\c "Figura" ',
             'Índice de figuras: se completa al actualizar los campos del documento (F9).')


def historial(w):
    w.titulo(1, 'Historial de revisión')
    w.p('Cambios del documento a lo largo del tiempo:')
    w.tabla((2.4, 1.8, 3.6, 7.7), ['Fecha', 'Versión', 'Autor', 'Descripción'], [
        [FECHA, VERSION, AUTOR, 'Versión inicial de N01 y N02: especificación de los procesos de negocio PN01 a PN04 '
                                 'y de los casos de uso CU-N01 a CU-N13, balanceada con el código y con la base de '
                                 'datos del sistema.']], centradas=(0, 1))


# ====================================================================== introduccion
def introduccion(w):
    w.titulo(1, 'Introducción')
    w.titulo(2, 'Propósito del documento')
    w.p('Este documento corresponde a la etapa N00 Procesos de negocio de la carpeta del proyecto. Contiene la '
        'especificación funcional de cada proceso de negocio de la Gestión Comercial de Quiniela (N01) y la '
        'especificación de cada caso de uso del sistema que los soporta (N02). Describe qué hace la agencia, cómo lo '
        'resuelve el sistema y cómo está construida cada funcionalidad, de manera que cada diagrama se pueda contrastar '
        'con el sistema en funcionamiento.')
    w.titulo(2, 'Alcance')
    w.p('Se especifican los cuatro procesos de negocio de la gestión comercial —PN01 Venta de apuestas, PN02 Gestión '
        'de sorteos y resultados, PN03 Pago de premios y PN04 Control de la gestión comercial— y los trece casos de uso '
        'del sistema que los soportan (CU-N01 a CU-N13). Las gestiones técnicas transversales (inicio y cierre de '
        'sesión, perfiles, idiomas, bitácora, control de cambios y dígitos verificadores) se documentan en los puntos '
        'T01 a T07; aquí solo se las nombra cuando intervienen, por ejemplo al registrar cada operación en la bitácora.')
    w.titulo(2, 'Fuentes y balanceo con el sistema')
    w.p('Los diagramas se construyeron a partir del código fuente de la solución —proyectos Entrega1software (UI), '
        'BLL, DAL y BE— y del script de la base de datos ScriptBD_AgenciaQuiniela.sql, de modo que quedan balanceados '
        'con la implementación:')
    w.vinetas([
        'Cada mensaje de un diagrama de secuencia es un método, constructor o propiedad que existe en la clase de su '
        'línea de vida, con la cantidad de parámetros que tiene en el código.',
        'Las clases, los atributos, las operaciones y las relaciones de los diagramas de clases existen en el código, '
        'y toda clase que recibe mensajes en un diagrama de secuencia figura en un diagrama de clases del mismo caso de '
        'uso, relacionada con la clase que le envía el mensaje.',
        'Las tablas que se leen o se escriben en los diagramas de secuencia están en el DER del caso de uso, con las '
        'columnas, claves y relaciones del script de la base de datos.',
        'Los textos de la interfaz que se citan entre comillas angulares («botones», «pestañas») y entre comillas '
        'dobles (“mensajes”) son exactamente los que muestra el sistema.',
        'Los prototipos de interfaz son capturas del sistema funcionando sobre la base de datos de prueba.',
    ])
    w.p('El control se hizo con un verificador automático que recorre todos los diagramas y todas las plantillas '
        '(Apéndice A). El resultado de la última verificación es de cero inconsistencias.')
    w.titulo(2, 'Convenciones de los diagramas')
    w.vinetas([
        '**Casos de uso (UML).** El actor principal se dibuja a la izquierda. El Administrador especializa al Operador '
        '(generalización), por lo que puede usar todos sus casos de uso; los que solo puede usar el Administrador '
        '(RN-17) se asocian directamente con él. Las relaciones «include» y «extend» se dibujan con flechas punteadas '
        'y los puntos de extensión con sus condiciones se indican en notas. En el modelo del negocio los procesos son '
        'casos de uso del negocio (elipse con una barra diagonal) y los actores llevan los estereotipos «actor del '
        'negocio» y «trabajador».',
        '**Actividad.** Un carril por rol. Las acciones que realiza o registra el sistema llevan el código del caso de '
        'uso del sistema que las soporta (ver la leyenda de cada diagrama).',
        '**Clases de diseño.** Paquetes por capa (UI, BLL, DAL y BE), visibilidad UML (+ público, - privado), '
        'estereotipos «static» y «singleton», asociaciones navegables con el nombre del atributo que las implementa '
        '(por ejemplo, -dal) y dependencias «create» y «use». Solo se muestran los miembros que intervienen en el caso '
        'de uso.',
        '**Secuencia (UML 2).** Los mensajes están numerados en orden. Los mensajes de retorno (línea punteada) se '
        'dibujan solo cuando el método devuelve un valor. Las líneas de vida sin «:» son clases estáticas. «create» '
        'indica la creación de un objeto y la X, su destrucción (fin de un bloque using o cierre de un formulario). '
        'Fragmentos combinados: alt (alternativas), opt (opcional), loop (repetición) y break (excepción que '
        'interrumpe el flujo, con el texto de la excepción). Un marco ref remite a otro diagrama del mismo caso de uso '
        'por su número. Las sentencias SQL se abrevian con «…».',
        '**Modelo conceptual.** Clases de análisis con sus atributos, sin identificadores ni operaciones, y '
        'asociaciones con nombre y multiplicidad.',
        '**DER.** Tablas con su clave primaria (PK) y sus claves foráneas (FK) y relaciones con su cardinalidad, tal '
        'como están en el script de la base de datos.',
        '**Diagramas de flujo.** Se incluyen solo cuando el algoritmo es complejo y complementa la descripción '
        'analítica.',
        '**Prototipos.** Capturas de pantalla con marcadores numerados en naranja, que se explican en la tabla que '
        'acompaña a cada figura.',
        '**Numeración.** Las figuras se numeran en todo el documento y su título empieza con el código del proceso o '
        'del caso de uso al que pertenecen. Cuando un caso de uso tiene varios diagramas del mismo tipo, el número '
        'entre paréntesis del marco del diagrama —por ejemplo, «sd CU-N01 (8)»— coincide con el del título de la '
        'figura («Diagrama de secuencia 8 de 12»).',
    ])
    w.titulo(2, 'Organización del documento')
    w.p('El capítulo N01 presenta el modelo de casos de uso del negocio, el modelo conceptual del dominio, los '
        'requisitos funcionales y las reglas de negocio, y luego cada proceso con sus cuatro puntos: identificación de '
        'roles, descripción funcional, diagrama de proceso y modelo conceptual. El capítulo N02 presenta el modelo de '
        'casos de uso del sistema y, para cada caso de uso, su especificación completa. Los archivos .drawio con los '
        'diagramas editables están en la carpeta Documentacion\\Diagramas del repositorio (Apéndice B).')


# ====================================================================== N01
def n01_general(w, figs):
    w.titulo(1, 'N01. Especificación funcional por proceso de negocio')
    w.p('En este capítulo se describe cada proceso de negocio de la Gestión Comercial de Quiniela con más detalle que '
        'en G02. Para cada proceso se identifican los roles intervinientes —con un diagrama de actores y un diagrama '
        'de secuencia de la colaboración—, se formaliza la descripción funcional Entrada / Comportamiento / Salida con '
        'sus requisitos funcionales y reglas de negocio, se modela el proceso con un diagrama de actividad y se '
        'presenta su modelo conceptual. Siguiendo la propuesta de Molina et al. (2000), “De los procesos del negocio a '
        'los casos de uso”, cada actividad del proceso que el sistema automatiza da origen a un caso de uso del '
        'sistema; la tabla de trazabilidad de cada proceso muestra esa correspondencia.')

    w.titulo(2, 'Modelo de casos de uso del negocio')
    w.p('El modelo muestra los cuatro procesos de negocio de la agencia como casos de uso del negocio, los actores del '
        'negocio —el Apostador y la Lotería— y los trabajadores —el Operador y el Administrador—. El ticket se imprime '
        'con la impresora predeterminada de Windows, que no se modela como actor: no tiene objetivos propios ni '
        'inicia ninguna interacción con el sistema. El Administrador es una especialización del Operador: interviene '
        'en todos los procesos del Operador y, además, cierra las ventas de los sorteos, registra los resultados y '
        'controla la gestión comercial con sus accesos directos.')
    w.figura(figs[('N01', 'uc_negocio')][0], max_alto=ALTO_TRAS_TITULO)
    roles = {}
    for pn, _ in PROCESOS:
        for nombre, tipo, desc, resp in n01.PROCESOS_N01[pn]['roles']:
            r = roles.setdefault(nombre, dict(tipo=tipo, desc=desc, procesos=[]))
            r['procesos'].append(pn)
    filas = []
    for nombre in ('Apostador', 'Lotería', 'Operador', 'Administrador'):
        r = roles[nombre]
        procesos = lista_y(r['procesos'])
        if nombre == 'Administrador':
            faltan = [pn for pn, _ in PROCESOS if pn not in r['procesos']]
            if faltan:
                procesos += '; en %s, como especialización del Operador' % lista_y(faltan)
        filas.append(['**%s**' % nombre, r['tipo'], r['desc'], procesos])
    w.p('**Actores y trabajadores del negocio**', antes=6, mantener=True)
    w.tabla((3.0, 3.2, 6.1, 3.2), ['Nombre', 'Tipo', 'Descripción', 'Procesos'], filas)
    w.p('**Procesos de negocio**', antes=12, mantener=True)
    w.tabla((1.8, 3.1, 7.2, 3.4), ['Código', 'Proceso', 'Objetivo', 'Casos de uso del sistema'],
            [[pn, nombre, n01.PROCESOS_N01[pn]['objetivo'], lista_y(n01.PROCESOS_N01[pn]['casos'])]
             for pn, nombre in PROCESOS], centradas=(0,))

    w.titulo(2, 'Modelo conceptual del dominio', nueva_pagina=True)
    w.p('El modelo conceptual reúne los conceptos de los cuatro procesos con sus atributos —sin identificadores ni '
        'operaciones— y las asociaciones entre ellos, con su nombre y su multiplicidad. El Premio no se guarda: se '
        'calcula al consultar o pagar un ticket. La asociación /es de entre Apuesta y Juego es derivada: el juego del '
        'ticket es el de los sorteos de sus jugadas.')
    w.figura(figs[('N01', 'conceptual')][0], max_alto=ALTO_TRAS_TITULO)
    w.p('**Conceptos del dominio**', antes=6, mantener=True)
    w.tabla((3.3, 7.8, 4.4), ['Concepto', 'Descripción', 'Atributos'],
            [['**%s**' % c, n01.DESCRIPCION_CONCEPTO[c], ', '.join(n01.ATRIBUTOS_CONCEPTO[c])]
             for c in n01.ATRIBUTOS_CONCEPTO])

    w.titulo(2, 'Requisitos funcionales', nueva_pagina=True)
    w.p('Requisitos funcionales de valor para el sistema que surgen de los procesos de negocio, con el proceso que los '
        'origina y los casos de uso que los implementan:')
    w.tabla((1.8, 8.2, 1.9, 3.6), ['Código', 'Requisito funcional', 'Proceso', 'Casos de uso'],
            [[rf, desc, proc, lista_y(casos)] for rf, desc, proc, casos in REQUISITOS], centradas=(0, 2))
    w.titulo(2, 'Reglas de negocio')
    w.p('Reglas de negocio que aplica el sistema; cada especificación de caso de uso indica las que le corresponden:')
    w.tabla((1.8, 13.7), ['Código', 'Regla de negocio'], [[rn, txt] for rn, txt in REGLAS], centradas=(0,))


def n01_proceso(w, figs, pn, nombre):
    p = n01.PROCESOS_N01[pn]
    w.titulo(2, '%s %s' % (pn, nombre), nueva_pagina=True)
    w.titulo(3, 'Objetivo del proceso')
    w.p(p['objetivo'])
    w.titulo(3, 'Descripción del proceso')
    for par in p['descripcion']:
        w.p(par)

    w.titulo(3, '1. Identificación de roles intervinientes')
    w.p('Roles que intervienen en el proceso, con su descripción y sus responsabilidades:', mantener=True)
    w.tabla((3.0, 3.0, 4.4, 5.1), ['Rol', 'Tipo', 'Descripción', 'Responsabilidades en el proceso'],
            [['**%s**' % n, t, d, r] for n, t, d, r in p['roles']])
    w.p('Diagrama de actores del proceso:', antes=12, mantener=True)
    w.figura(figs[(pn, 'actores')][0], max_alto=ALTO_TRAS_TITULO)
    w.p('La colaboración entre los roles a lo largo del proceso se modela con un diagrama de secuencia de nivel de '
        'negocio, en el que el sistema se representa con una única línea de vida:', mantener=True)
    w.figura(figs[(pn, 'colab')][0], max_alto=ALTO_TRAS_TITULO)

    w.titulo(3, '2. Descripción funcional del proceso: Entrada / Comportamiento / Salida', nueva_pagina=True)
    t = w.tabla((3.4, 12.1), ['Elemento', 'Descripción'])
    for etiqueta, clave in (('Entrada', 'entradas'), ('Comportamiento', 'comportamiento'), ('Salida', 'salidas')):
        w.fila(t, [etiqueta, p[clave]], negritas=[True, False], fondos=[AZUL_CLARO, None], vinetas=True,
               no_dividir=False)
    w.p('**Requisitos funcionales que formaliza el proceso**', antes=12, mantener=True)
    reqs = {rf: (desc, casos) for rf, desc, proc, casos in REQUISITOS}
    w.tabla((1.8, 10.1, 3.6), ['Código', 'Requisito funcional', 'Casos de uso'],
            [[rf, reqs[rf][0], lista_y(reqs[rf][1])] for rf in p['requisitos']], centradas=(0,))
    w.p('**Reglas de negocio que aplica**', antes=12, mantener=True)
    w.tabla((1.8, 13.7), ['Código', 'Regla de negocio'], [[rn, TEXTO_REGLA[rn]] for rn in p['reglas']],
            centradas=(0,))

    w.titulo(3, '3. Diagrama de proceso', nueva_pagina=True)
    acts = figs[(pn, 'actividad')]
    w.p('El proceso se modela con %s, con un carril por rol. Las acciones que realiza o registra el sistema '
        'llevan el código del caso de uso del sistema que las soporta.'
        % ('un diagrama de actividad' if len(acts) == 1 else '%d diagramas de actividad' % len(acts)), mantener=True)
    for k, f in enumerate(acts):
        w.figura(f, max_alto=ALTO_TRAS_TITULO if k == 0 else 19.0)
    w.p('**Del proceso a los casos de uso del sistema**', antes=12, mantener=True)
    w.p('Actividades del proceso que automatiza el sistema y casos de uso que las soportan:', mantener=True)
    w.tabla((7.5, 8.0), ['Actividad del proceso', 'Casos de uso del sistema'],
            [[act, [cu_nombre(c) for c in casos]] for act, casos in n01.TRAZABILIDAD[pn]])

    w.titulo(3, '4. Modelo conceptual', nueva_pagina=True)
    w.p('Conceptos del dominio que intervienen en el proceso: las entidades que surgen de los flujos de información '
        'del diagrama de actividad.', mantener=True)
    w.figura(figs[(pn, 'conceptual')][0], max_alto=ALTO_TRAS_TITULO)
    w.tabla((3.3, 12.2), ['Concepto', 'Descripción'],
            [['**%s**' % c, n01.DESCRIPCION_CONCEPTO[c]] for c in p['conceptos']])


def n01_trazabilidad(w):
    w.titulo(2, 'Trazabilidad entre requisitos funcionales y casos de uso', nueva_pagina=True)
    w.p('Cada requisito funcional está implementado por al menos un caso de uso y cada caso de uso implementa al menos '
        'un requisito:', mantener=True)
    casos = list(CASOS)
    anchos = [1.6] + [round(13.9 / len(casos), 2)] * len(casos)
    t = w.tabla(anchos, ['RF'] + [c.replace('CU-', '') for c in casos], tam=8, centradas=tuple(range(len(anchos))))
    for rf, desc, proc, cs in REQUISITOS:
        w.fila(t, [rf] + ['X' if c in cs else '' for c in casos], tam=8.5, centradas=tuple(range(len(anchos))),
               negritas=[True] + [False] * len(casos))
    w.p('Las columnas N01 a N13 corresponden a los casos de uso CU-N01 a CU-N13.', tam=9, antes=6)


# ====================================================================== N02
def relaciones(spec):
    c = spec['codigo']
    res = []
    for base, punto, cond in spec.get('extiende', []):
        res.append('Extiende %s (punto %s)' % (base, punto))
    for punto, paso, ext, cond in spec.get('extension', []):
        res.append('Extendido por %s (punto %s, paso %d)' % (ext, punto, paso))
    for item in spec.get('inclusion', []):
        res.append('Incluye %s (paso %d)' % (item[0], item[1]))
    if spec.get('incluido_por'):
        res.append('Incluido por %s' % lista_y(spec['incluido_por']))
    return res or ['—']


def n02_general(w, figs, specs):
    w.titulo(1, 'N02. Especificaciones de casos de uso')
    w.p('Para cada caso de uso identificado en N01 se presenta una especificación, organizada por proceso de negocio. '
        'Cada especificación empieza con el gráfico del caso de uso (ítem h de la plantilla) y, debajo, la plantilla '
        'con la carátula (a), el historial de revisión (b), el objetivo (c), las precondiciones (d), las '
        'postcondiciones (e), los eventos disparadores (f), los puntos de extensión y condiciones (g) y la descripción '
        'analítica con el flujo normal y los flujos alternativos (i). A continuación, cada uno en una nueva carilla, '
        'siguen el diagrama de flujo (j, solo cuando el algoritmo es complejo), los diagramas de clases (k), los '
        'diagramas de secuencia (l), el DER (m) y los prototipos de interfaz de usuario (n).')
    w.titulo(2, 'Modelo de casos de uso del sistema')
    w.p('El modelo reúne los trece casos de uso de la gestión comercial. El Operador usa los casos de venta, de '
        'post-venta, de pago y de consulta; el Administrador lo especializa, de modo que puede usar todos ellos y, '
        'además, los que solo están habilitados para su perfil (RN-17): cerrar las ventas, cargar el resultado y '
        'sortear al azar. Estos tres son objetivos propios del Administrador, por eso se asocian directamente con él '
        'y no extienden la consulta de números ganadores, aunque se inicien desde la misma pantalla. CU-N03 '
        'Imprimir ticket extiende las dos ventas y la reimpresión: el ticket siempre se muestra en pantalla y se '
        'imprime solo si el operador presiona «Imprimir» (condición de la extensión). CU-N04 Buscar ticket está '
        'incluido en la reimpresión, la cancelación y el pago, que no pueden hacerse sin encontrar antes el ticket. '
        'La impresora no se modela como actor: es la impresora predeterminada de Windows, que el sistema usa por '
        'medio de PrintDocument y que no inicia ninguna interacción.')
    w.figura(figs[('N02', 'uc_sistema')][0], max_alto=ALTO_TRAS_TITULO)
    w.p('**Catálogo de casos de uso**', antes=6, mantener=True)
    w.tabla((1.9, 3.2, 1.8, 2.6, 3.8, 2.2),
            ['Código', 'Caso de uso', 'Proceso', 'Actor principal', 'Relaciones', 'Requisitos'],
            [[s['codigo'], CASOS[s['codigo']][0], CASOS[s['codigo']][1], CASOS[s['codigo']][2], relaciones(s),
              lista_y(requisitos_de(s['codigo']))] for s in specs], tam=9, centradas=(0, 2))


def plantilla(w, spec, n_fig_uc):
    c = spec['codigo']
    nombre, proceso, actor = CASOS[c]
    t = w.tabla(ANCHOS_PLANTILLA)
    todo = [(0, 4)]
    rotulo = [(0, 1), (2, 4)]
    tres = [(0, 1), (2, 2), (3, 4)]

    def barra(texto):
        w.fila(t, [texto], unir=todo, fondo=BARRA, negrita=True, color=AZUL, mantener=True)

    def campo(etiqueta, valor, vinetas=None):
        w.fila(t, [etiqueta, valor], unir=rotulo, fondos=[AZUL_CLARO, None], negritas=[True, False], vinetas=vinetas)

    enc = w.fila(t, ['ESPECIFICACIÓN DE CASO DE USO %s  ·  %s' % (esp(c), cu_nombre(c))], unir=todo, fondo=AZUL,
                 negrita=True, color='FFFFFF', tam=10)
    # el titulo de la plantilla se repite arriba de cada pagina que ocupa
    enc._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))

    barra('a. Carátula')
    campo('Código de la especificación', esp(c))
    campo('Caso de uso', cu_nombre(c))
    campo('Proceso de negocio', '%s %s' % (proceso, NOMBRE_PROCESO[proceso]))
    campo('Actor principal', spec['actor'])
    campo('Actores secundarios', spec['secundarios'])
    campo('Requisitos funcionales', lista_y(requisitos_de(c)))
    campo('Reglas de negocio', lista_y(spec['reglas']))

    barra('b. Historial de revisión')
    historial = [(0, 1), (2, 2), (3, 3), (4, 4)]
    w.fila(t, ['Fecha', 'Versión', 'Autor', 'Descripción'], unir=historial, fondo=AZUL_CLARO, negrita=True,
           mantener=True)
    w.fila(t, [FECHA, VERSION, AUTOR, 'Versión inicial de la especificación, balanceada con el código y con la base '
                                      'de datos del sistema.'], unir=historial, centradas=(0, 1))

    barra('c. Objetivo')
    w.fila(t, [spec['objetivo']], unir=todo)
    barra('d. Precondiciones')
    w.fila(t, [spec['pre']], unir=todo, vinetas=True)
    barra('e. Postcondiciones')
    campo('Con éxito', spec['post'], vinetas=True)
    campo('Sin éxito', spec['post_fracaso'], vinetas=True)
    barra('f. Eventos disparadores')
    w.fila(t, [spec['disparadores']], unir=todo, vinetas=True)

    barra('g. Puntos de extensión y condiciones')
    hubo = False
    if spec.get('extension'):
        hubo = True
        w.fila(t, ['Punto de extensión', 'Caso de uso que extiende', 'Condición'], unir=tres,
               fondo=AZUL_CLARO, negrita=True, mantener=True)
        for punto, paso, ext, cond in spec['extension']:
            w.fila(t, ['%s (paso %d)' % (punto, paso), ext, cond], unir=tres)
    if spec.get('extiende'):
        hubo = True
        w.fila(t, ['Caso de uso base', 'Punto de extensión', 'Condición'], unir=tres,
               fondo=AZUL_CLARO, negrita=True, mantener=True)
        for base, punto, cond in spec['extiende']:
            w.fila(t, [cu_nombre(base), punto, cond], unir=tres)
    if spec.get('inclusion'):
        hubo = True
        w.fila(t, ['Caso de uso incluido', 'Paso', 'Descripción'], unir=tres, fondo=AZUL_CLARO,
               negrita=True, mantener=True)
        for inc, paso, desc in spec['inclusion']:
            w.fila(t, [cu_nombre(inc), str(paso), desc], unir=tres, centradas=(1,))
    if spec.get('incluido_por'):
        hubo = True
        w.fila(t, ['Este caso de uso no tiene puntos de extensión: está incluido en %s, que lo usan para buscar el '
                   'ticket.' % lista_y(cu_nombre(x) for x in spec['incluido_por'])], unir=todo)
    if not hubo:
        w.fila(t, ['El caso de uso no tiene puntos de extensión ni inclusiones.'], unir=todo)

    barra('h. Gráfico del caso de uso')
    fila = w.fila(t, [''], unir=todo)
    par = fila.cells[0].paragraphs[0]
    w.referencia(par, n_fig_uc)
    r = par.add_run(', al comienzo de esta especificación: el caso de uso con sus actores y sus relaciones.')
    r.font.size = Pt(9.5)
    for run in par.runs:
        run.font.size = Pt(9.5)

    barra('i. Descripción analítica: flujo normal')
    w.fila(t, ['N°', 'Actor', 'Acción'], unir=[(0, 0), (1, 1), (2, 4)], fondo=AZUL_CLARO, negrita=True,
           centradas=(0,), mantener=True)
    for k, (quien, texto) in enumerate(spec['flujo'], start=1):
        w.fila(t, [str(k), actor if quien == 'A' else 'Sistema', texto], unir=[(0, 0), (1, 1), (2, 4)],
               centradas=(0,))
    barra('i. Descripción analítica: flujos alternativos')
    w.fila(t, ['Flujo', 'Descripción'], unir=[(0, 0), (1, 4)], fondo=AZUL_CLARO, negrita=True, centradas=(0,),
           mantener=True)
    for ref, tit, lineas, ret in spec['alternativos']:
        cuerpo = ['**%s.**' % tit.rstrip('.')]
        if len(lineas) > 1:
            cuerpo += ['%d. %s' % (i, l) for i, l in enumerate(lineas, start=1)]
        else:
            cuerpo += list(lineas)
        cuerpo.append('**Retorno:** %s' % ret)
        w.fila(t, [ref, cuerpo], unir=[(0, 0), (1, 4)], centradas=(0,), vinetas=False)
    return t


def n02_caso(w, figs, spec, flujos):
    c = spec['codigo']
    w.titulo(2, cu_nombre(c), nueva_pagina=True)
    n_uc = w.figura(figs[(c, 'uc')][0], max_alto=10.0, rotar=False)
    w.titulo(3, 'a – i. Plantilla de especificación %s' % esp(c))
    plantilla(w, spec, n_uc)

    # j. diagrama de flujo
    if spec.get('flujo_alg'):
        w.titulo(3, 'j. Diagrama de flujo', nueva_pagina=True)
        w.p(spec['flujo_alg'][0], mantener=True)
        f = figs[(c, 'flujo')][0]
        flujos[c] = w.figura(f, max_alto=ALTO_TRAS_TITULO)
    else:
        w.titulo(3, 'j. Diagrama de flujo')
        texto = spec.get('flujo_no_aplica', 'No corresponde: el caso de uso no tiene un algoritmo complejo.')
        if not texto.startswith('No corresponde'):
            texto = 'No corresponde un diagrama propio. ' + texto
        par = w.p(texto)
        otro = {'CU-N02': 'CU-N01', 'CU-N10': 'CU-N09'}.get(c)
        if otro and otro in flujos:
            par.add_run(' Ver ')
            w.referencia(par, flujos[otro])
            par.add_run('.')

    # k. clases
    w.titulo(3, 'k. Diagramas de clases', nueva_pagina=True)
    cls = figs[(c, 'clases')]
    w.p('Clases de diseño que intervienen en el caso de uso, agrupadas por capa (UI, BLL, DAL y BE), con los atributos '
        'y las operaciones del código que se usan en sus diagramas de secuencia%s.'
        % ('' if len(cls) == 1 else ' (%d diagramas)' % len(cls)), mantener=True)
    for k, f in enumerate(cls):
        w.figura(f, max_alto=ALTO_TRAS_TITULO if k == 0 else 19.0)

    # l. secuencia
    w.titulo(3, 'l. Diagramas de secuencia', nueva_pagina=True)
    sds = figs[(c, 'sd')]
    w.p('El caso de uso se modela con %s. El título de cada figura indica los pasos y los flujos alternativos de la '
        'descripción analítica que cubre; el número entre paréntesis del marco coincide con el del título y es el que '
        'usan los marcos ref para remitir de un diagrama a otro.'
        % ('un diagrama de secuencia' if len(sds) == 1 else '%d diagramas de secuencia' % len(sds)), mantener=True)
    for k, f in enumerate(sds):
        w.figura(f, max_alto=ALTO_TRAS_TITULO if k == 0 else 19.0)

    # m. DER
    w.titulo(3, 'm. DER', nueva_pagina=True)
    # el DER va directamente debajo del titulo para que tenga casi toda la hoja; la explicacion va despues
    w.figura(figs[(c, 'der')][0], max_alto=18.3)
    w.p('El DER muestra las tablas de la base de datos que lee o escribe el caso de uso, con sus claves primarias y '
        'foráneas y la cardinalidad de cada relación.')
    if spec.get('der_nota'):
        w.p(spec['der_nota'])
    w.p('**Correspondencia entre las clases de entidad y las tablas**', antes=10, mantener=True)
    w.tabla((3.3, 3.3, 8.9), ['Clase (BE)', 'Tabla', 'Observaciones'],
            [[cl, CORRESPONDENCIA[cl][0], CORRESPONDENCIA[cl][1]] for cl in spec['correspondencia']])

    # n. prototipos
    w.titulo(3, 'n. Prototipos de interfaz de usuario', nueva_pagina=True)
    protos = figs.get((c, 'proto'), [])
    w.p('Capturas del sistema funcionando sobre la base de datos de prueba. Los números en naranja identifican los '
        'elementos que se describen en la tabla de cada figura, con el paso o el flujo de la descripción analítica en '
        'que se usan.', mantener=True)
    for f in protos:
        w.figura(f)
        w.tabla((0.9, 3.3, 2.6, 6.5, 2.2), ['N°', 'Elemento', 'Tipo', 'Descripción', 'Paso'], f.filas, tam=9,
                centradas=(0, 4))
        if f.nota:
            w.p(f.nota, tam=9.5, antes=6)


# ====================================================================== apendices
def apendice_balanceo(w, specs, figs):
    w.titulo(1, 'Apéndice A. Verificación automática del balanceo')
    w.p('Para que la documentación quede balanceada con la implementación —criterio de evaluación de la cátedra para '
        'los casos de uso, los diagramas de secuencia, los diagramas de clases y el DER—, los diagramas de este '
        'documento se generan con un programa que toma los datos del código fuente y del script de la base de datos '
        '(carpeta Documentacion\\generador del repositorio), y un verificador (balanceo.py) controla cada caso de uso '
        'antes de armar el documento. Para cada caso de uso se verifica que:')
    w.vinetas([
        'cada mensaje sincrónico o de creación de un diagrama de secuencia sea un método, un constructor o una '
        'propiedad que existe en la clase de la línea de vida, con una cantidad de parámetros válida;',
        'cada clase que recibe mensajes y cada operación que recibe figuren en un diagrama de clases del caso de uso, '
        'y que entre las clases que se envían mensajes haya una relación dibujada;',
        'cada tabla que se lee o se escribe en los mensajes a la base de datos esté en el DER del caso de uso;',
        'los textos citados entre comillas (mensajes del sistema) existan en el código y los botones, pestañas y '
        'etiquetas citados entre comillas angulares existan en los formularios;',
        'las referencias a pasos, flujos alternativos, puntos de extensión y reglas de negocio sean coherentes con la '
        'descripción analítica;',
        'cada requisito funcional esté implementado por al menos un caso de uso y cada caso de uso implemente al menos '
        'un requisito.',
    ])
    w.p('Resultado de la última verificación (%s):' % FECHA, antes=8, mantener=True)
    filas = []
    tot = dict(sd=0, cl=0, mensajes=0, tablas=0, textos=0, prob=0)
    for s in specs:
        st = {}
        prob = [x for x in verificar(s, st) if not x.startswith('(info)')]
        c = s['codigo']
        nsd, ncl = len(s['secuencias']), len(s['clases'])
        textos = st['citas'] + st['etiquetas']
        filas.append([c, str(nsd), str(st['mensajes']), str(ncl), str(st['tablas']), str(textos), str(len(prob))])
        for k, v in (('sd', nsd), ('cl', ncl), ('mensajes', st['mensajes']), ('tablas', st['tablas']),
                     ('textos', textos), ('prob', len(prob))):
            tot[k] += v
    t = w.tabla((2.0, 2.2, 2.3, 2.1, 1.9, 2.1, 2.9),
                ['Caso de uso', 'Diagramas de secuencia', 'Mensajes verificados', 'Diagramas de clases',
                 'Tablas usadas', 'Textos y etiquetas', 'Inconsistencias'], filas, centradas=tuple(range(7)))
    w.fila(t, ['Total', str(tot['sd']), str(tot['mensajes']), str(tot['cl']), str(tot['tablas']), str(tot['textos']),
               str(tot['prob'])], negrita=True, fondo=AZUL_CLARO, centradas=tuple(range(7)))
    return tot


def apendice_archivos(w, figs_todas):
    w.titulo(1, 'Apéndice B. Archivos de los diagramas')
    w.p('Los diagramas se entregan también en formato editable de draw.io (diagrams.net), en la carpeta '
        'Documentacion\\Diagramas del repositorio: un archivo por grupo y, dentro de cada archivo, una página por '
        'diagrama con el mismo título que su marco. Los prototipos son capturas del sistema y no tienen archivo .drawio.')
    grupos = {}
    for f in figs_todas:
        if f.tipo != 'proto':
            grupos.setdefault(catalogo.nombre_archivo(f.grupo), []).append(f)
    filas = []
    for archivo, fs in grupos.items():
        tipos = {}
        for f in fs:
            tipos[f.tipo] = tipos.get(f.tipo, 0) + 1
        nombres = dict(uc_negocio='casos de uso del negocio', conceptual='modelo conceptual', actores='actores',
                       colab='colaboración', actividad='actividad', uc_sistema='casos de uso del sistema',
                       uc='casos de uso', flujo='flujo', clases='clases', sd='secuencia', der='DER')
        detalle = ', '.join('%d de %s' % (n, nombres[k]) for k, n in tipos.items())
        filas.append([archivo, str(len(fs)), detalle])
    w.tabla((6.4, 1.9, 7.2), ['Archivo', 'Páginas', 'Diagramas'], filas, centradas=(1,))


# ====================================================================== armado
def armar(salida=SALIDA):
    figs_todas = catalogo.todas(con_prototipos=True)
    figs = {}
    for f in figs_todas:
        figs.setdefault((f.grupo, f.tipo), []).append(f)
    faltan = [f.ruta for f in figs_todas if not Path(f.ruta).exists()]
    if faltan:
        raise SystemExit('Faltan imagenes (correr catalogo.py): %s' % faltan[:5])
    specs = especificaciones()

    w = Word(REF, LOGO, ROTADAS)
    pendientes = w.encabezado('rId9', {
        '<Nombre del Sistema>': SISTEMA_CORTO,
        '<Nombre de Sección / Etapa> 1 Entrega': 'N01 y N02 - 1 Entrega',
        'Fecha 8/9/2026': 'Fecha %s' % FECHA,
        # "Versión" queda como en la plantilla: el número no entra en su cuadro (la versión está en la portada
        # y en el historial de revisión)
    }, achicar={'Fecha 8/9/2026'})
    if pendientes:
        print('AVISO: textos del encabezado que no se encontraron:', pendientes)
    w.pie_centrado('rId10')
    cp = w.doc.core_properties
    cp.title = '%s - N01 y N02' % SISTEMA_CORTO
    cp.subject = 'N01 Especificación funcional por proceso de negocio y N02 Especificaciones de casos de uso'
    cp.author = AUTOR
    cp.last_modified_by = AUTOR
    cp.keywords = 'Trabajo de Diploma; UAI; Agencia de Quiniela; casos de uso'
    cp.comments = ''
    cp.revision = 1

    portada(w)
    indices(w)
    historial(w)
    introduccion(w)
    n01_general(w, figs)
    for pn, nombre in PROCESOS:
        n01_proceso(w, figs, pn, nombre)
    n01_trazabilidad(w)
    n02_general(w, figs, specs)
    flujos = {}
    for s in specs:
        n02_caso(w, figs, s, flujos)
    tot = apendice_balanceo(w, specs, figs)
    apendice_archivos(w, figs_todas)
    w.seccion_final(w.seccion(True, 'rId9', 'rId10', inicio=2))
    quitadas = w.limpiar_relaciones()
    w.guardar(salida)
    return w, tot, quitadas


if __name__ == '__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    salida = Path(sys.argv[1]) if len(sys.argv) > 1 else SALIDA
    w, tot, quitadas = armar(salida)
    print('figuras:', len(w.figuras), '| relaciones quitadas:', len(quitadas), '| inconsistencias:', tot['prob'])
    chicas = [(n, f.archivo, round(pt, 1), g) for n, f, pt, g in w.figuras if pt is not None and pt < 6]
    print('figuras con letra menor a 6 pt:', len(chicas))
    for x in chicas:
        print('   ', x)
    print('giradas:', sum(1 for x in w.figuras if x[3]))
    print('guardado en', salida)
